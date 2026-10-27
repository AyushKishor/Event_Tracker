#!/usr/bin/env python3
"""Replay a manually-kept tracker (xlsx) through the filters and list what would be skipped.

Usage: python scripts/coverage_check.py path/to/Events_Tracker.xlsx
Each row is tested twice: name + city only ("name only": what a scraped page gives with no
account label), and with the Target Account added as host ("with account").
"""
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent.parent))
from refresh import load_config  # noqa: E402
from tracker import extract  # noqa: E402

cfg = load_config(Path(__file__).parent.parent / "config.yaml")
H = ["Target Account", "Event name", "City", "Country", "Virtual/In-person"]


def rows(path):
    wb = load_workbook(path, data_only=True)
    for ws in wb:
        hdr = [c.value for c in ws[1]]
        if "Event name" not in hdr:
            continue
        ix = {h: hdr.index(h) for h in H if h in hdr}
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r[ix["Event name"]]:
                yield ws.title, {h: (r[i] or "") for h, i in ix.items()}


def verdict(row, with_account):
    acct = str(row["Target Account"])
    host = acct if with_account and acct not in ("-", "") else ""
    fmt = "Virtual" if str(row["Virtual/In-person"]).lower() == "virtual" else ""
    return extract.classify(" ".join(str(row["Event name"]).split()), host, "", str(row["City"]).replace("-", ""),
                            fmt, {}, cfg)


total = miss_name = miss_acct = 0
for sheet, row in rows(sys.argv[1]):
    total += 1
    a, b = verdict(row, False), verdict(row, True)
    miss_name += not a.keep
    miss_acct += not b.keep
    if not b.keep or not a.keep:
        tag = "SKIP " if not b.keep else "name-only skip"
        print(f"{tag:15} [{sheet}] {row['Target Account']!s:22} {' '.join(str(row['Event name']).split())[:60]:60} {row['City']}")
print(f"\n{total} events: skipped {miss_name} on name only, {miss_acct} with account label")
