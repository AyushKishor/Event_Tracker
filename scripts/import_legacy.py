#!/usr/bin/env python3
"""One-off: load the manually-kept tracker (xlsx) into tracker.xlsx so history is deduped.

Usage: python scripts/import_legacy.py path/to/Events_Tracker_2026.xlsx [--output tracker.xlsx]
"""
import argparse
import re
import sys
from datetime import date, datetime
from pathlib import Path

from dateutil import parser as dateparser
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent.parent))
from refresh import load_config  # noqa: E402
from tracker import extract, workbook  # noqa: E402

ROOT = Path(__file__).parent.parent
FMT = {"in-person": "In-person", "virtual": "Virtual", "both": "Hybrid"}


def clean(v):
    s = " ".join(str(v).split()) if v is not None else ""
    return "" if s in ("-", "TBD", "None") else s


def parse_dates(day, month, year):
    """'11-13', 'January', 2026 -> (2026-01-11, 2026-01-13); a datetime or '3–7 Aug 2026' also work."""
    if isinstance(day, datetime):
        return day.date(), None
    text = str(day or "").replace("–", "-").strip()
    try:
        if re.search(r"[A-Za-z]", text):  # '3-7 Aug 2026'
            m = re.match(r"(\d{1,2})(?:-(\d{1,2}))?\s+(.+)", text)
            d1, d2, rest = m.groups()
            return (dateparser.parse(f"{d1} {rest}").date(),
                    dateparser.parse(f"{d2} {rest}").date() if d2 else None)
        if not (month and year and text):
            return None, None
        nums = [int(n) for n in re.findall(r"\d+", text)]
        start = dateparser.parse(f"{nums[0]} {month} {year}").date()
        end = dateparser.parse(f"{nums[1]} {month} {year}").date() if len(nums) > 1 and nums[1] >= nums[0] else None
        return start, end
    except (ValueError, AttributeError, IndexError, OverflowError):
        return None, None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("legacy", type=Path)
    ap.add_argument("--output", type=Path, default=ROOT / "tracker.xlsx")
    args = ap.parse_args()
    cfg = load_config(ROOT / "config.yaml")

    events, skipped, dropped = [], [], []
    for ws in load_workbook(args.legacy, data_only=True):
        hdr = [c.value for c in ws[1]]
        if "Event name" not in hdr:
            continue
        for r in ws.iter_rows(min_row=2, values_only=True):
            row = dict(zip(hdr, r))
            name = clean(row.get("Event name"))
            if not name:
                continue
            start, end = parse_dates(row.get("Date"), row.get("Month"), row.get("Year"))
            if not start:
                skipped.append(name)
                continue
            city = clean(row.get("City"))
            acct = clean(row.get("Target Account"))
            v = extract.classify(name, acct if acct != "-" else "", "", city, "", {}, cfg)
            if not v.keep:  # same relevance rules as the scraper: no small community meetups
                dropped.append(name)
                continue
            found = extract.detect_companies(f"{name} {acct}", cfg["_companies"])
            link = clean(row.get("Link"))
            events.append(extract.Event(
                name=name, start=start, end=end, city=city,
                host=v.host or (acct if acct not in ("-", "") and "/" not in acct else ""),
                region=extract.detect_region(city, cfg["_regions"]) or "Other",
                fmt=FMT.get(clean(row.get("Virtual/In-person")).lower(), ""),
                companies=found or ([acct] if acct and "/" not in acct else []),
                cost=clean(row.get("Entry Fee (USD)")),
                url=link if link.startswith("http") else "",
                source="Legacy tracker 2026",
            ))
    added, total = workbook.merge(args.output, events, earliest=cfg.get("earliest_date"))
    print(f"Imported {added} new event(s); {total} total. Skipped (no usable date): {skipped}")
    print(f"Dropped as small/not company-organised ({len(dropped)}): {dropped}")


if __name__ == "__main__":
    main()
