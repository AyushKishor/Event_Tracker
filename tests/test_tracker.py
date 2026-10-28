import json
import sys
from datetime import date
from pathlib import Path

import yaml

sys.path.insert(0, str(Path(__file__).parent.parent))

from refresh import load_config  # noqa: E402
from tracker import extract, workbook  # noqa: E402

ROOT = Path(__file__).parent.parent
TODAY = date(2027, 1, 15)


def cfg(keep_out_of_region=True):
    """Test config. Most tests exercise tagging with out-of-region events kept; the shipped default is strict."""
    c = load_config(ROOT / "config.yaml")
    c["keep_out_of_region"] = keep_out_of_region
    return c


def jsonld_page(events):
    return "<html><script type='application/ld+json'>" + json.dumps(events) + "</script></html>"


def ev(name, start, city, organizer="", mode=None, performer=None):
    d = {"@type": "Event", "name": name, "startDate": start,
         "location": {"@type": "Place", "address": {"addressLocality": city}}}
    if organizer:
        d["organizer"] = {"name": organizer}
    if mode:
        d["eventAttendanceMode"] = mode
    if performer:
        d["performer"] = [{"name": p} for p in performer]
    return d


def test_company_alias_matching():
    c = cfg()["_companies"]
    assert extract.detect_companies("AWS and Microsoft partner day", c) == ["Microsoft", "AWS"]
    assert extract.detect_companies("laws of physics", c) == []  # no substring match on "aws"
    assert extract.detect_companies("Oracle Cloud Infrastructure", c) == ["Oracle"]


def test_region_detection():
    r = cfg()["_regions"]
    assert extract.detect_region("Austin", r) == "Texas"
    assert extract.detect_region("Bengaluru, Karnataka", r) == "India"
    assert extract.detect_region("Redmond", r) == "Seattle"
    assert extract.detect_region("Chicago", r) == ""
    assert extract.detect_region("Brooklyn, NY", r) == "New York"


def test_parse_filters_region_window_and_company():
    html = jsonld_page([
        ev("AWS Partner Summit", "2027-03-10", "Seattle", "AWS", performer=["Jane Doe"]),
        ev("Random Meetup", "2027-03-10", "Seattle", "Someone"),             # no tracked company
        ev("Google Cloud Partner Day", "2027-03-11", "Chicago", "Google"),   # wrong region
        ev("Oracle Partner Webinar", "2027-03-12", "", "Oracle",
           mode="https://schema.org/OnlineEventAttendanceMode"),            # virtual kept
        ev("Salesforce Partner Day", "2029-01-01", "Austin", "Salesforce"),  # outside window
        ev("Past Partner Day", "2026-01-01", "Austin", "Salesforce"),        # past
    ])
    out = extract.parse_events(html, "http://x", cfg(), {"label": "t", "require_company": True}, TODAY)
    assert [e.name for e in out] == ["AWS Partner Summit", "Google Cloud Partner Day", "Oracle Partner Webinar"]
    aws, google, oracle = out
    assert (google.region, google.companies) == ("Other", ["Google"])  # out-of-region flagship-style kept
    assert (aws.region, aws.companies, aws.role, aws.speakers) == ("Seattle", ["AWS"], "Host", "Jane Doe")
    assert oracle.region == "Virtual"


def test_virtual_toggle():
    c = cfg()
    c["include_virtual"] = False
    c["keep_out_of_region"] = False
    html = jsonld_page([ev("Oracle Webinar", "2027-03-12", "", "Oracle", mode="OnlineEventAttendanceMode")])
    assert extract.parse_events(html, "http://x", c, {"label": "t"}, TODAY) == []


def test_card_fallback_and_date_range():
    html = """<ul><li><a href="/e/1"><h3>ServiceNow Partner Forum</h3></a>
              <span>Mar 3-5, 2027</span> Austin, TX</li></ul>"""
    out = extract.parse_events(html, "http://x.com/events", cfg(), {"label": "t", "host": "ServiceNow"}, TODAY)
    assert len(out) == 1
    e = out[0]
    assert (e.start, e.end, e.region, e.url) == (date(2027, 3, 3), date(2027, 3, 5), "Texas", "http://x.com/e/1")
    assert e.companies == ["ServiceNow"] and e.role == "Host"


def test_dedupe_key_normalises():
    a = workbook.event_key("AWS  Partner Summit!", date(2027, 3, 10), "Seattle")
    b = workbook.event_key("aws partner summit", date(2027, 3, 10), " seattle ")
    assert a == b
    assert a != workbook.event_key("aws partner summit", date(2027, 3, 11), "seattle")


def test_merge_idempotent_single_sheet_no_tracking_columns(tmp_path):
    path = tmp_path / "t.xlsx"
    html = jsonld_page([ev("AWS Partner Summit", "2027-03-10", "Seattle", "AWS")])
    events = extract.parse_events(html, "http://x", cfg(), {"label": "t"}, TODAY)

    assert workbook.merge(path, events, TODAY) == (1, 1)
    assert workbook.merge(path, events, TODAY) == (0, 1)  # no duplicate

    from openpyxl import load_workbook
    wb = load_workbook(path)
    assert wb.sheetnames == ["2026", "2027"]
    assert [c.value for c in wb["2027"][1]] == workbook.HEADERS
    assert not {"Date Added", "Status", "Notes"} & set(workbook.HEADERS)


def verdict(name, city="", host="", fmt=""):
    return extract.classify(name, host, "", city, fmt, {}, cfg())


def test_events_from_2026_sheet_are_covered():
    # (name, city, host) rows from the manually kept tracker that a narrow filter would drop
    keep = [
        ("Channel Partners Conference & Expo", "Las Vegas", ""),     # channel industry, out of region
        ("PartnerTechX", "Phoenix", ""),
        ("Catalyst by Partnership Leaders", "", ""),
        ("CES 2026", "Las Vegas", ""),                                # general tech, always_track
        ("TechCrunch Disrupt 2026", "San Francisco", ""),
        ("Knowledge 2026", "Las Vegas", ""),                          # vendor flagship, out of region
        ("Dreamforce 2026", "San Francisco", ""),
        ("Snowflake Summit 2026", "Moscone Center San Francisco", ""),  # non-core target account
        ("Dell Technologies World", "Las Vegas", ""),
        ("Microsoft AI Tour", "Boston", ""),
        ("AWS Marketplace Seller Forum", "Bellevue", ""),
        ("Oracle AI World Tour", "Mumbai", ""),
        ("Cypher 2026", "Bengaluru", ""),
        ("Build with Gemini - Seattle", "Seattle", ""),               # Google-run event
        ("SaaSify Partner Central Agentic Workshop", "Seattle", "AWS"),
    ]
    for name, city, host in keep:
        assert verdict(name, city, host).keep, name


def test_irrelevant_out_of_region_event_is_dropped():
    assert not verdict("Local Pottery Fair", "Boise").keep
    assert not verdict("Regional Dental Expo", "Boise").keep  # event word alone, out of region


SERVICENOW_MD = """
World Forum London, United Kingdom

Oct 20, 2026Excel London
Video CTA

![tokyo city image](https://example.com/x.png)World Forum Tokyo, Japan
Oct 27, 2026 - Oct 28, 2026Grand Prince Hotel Takanawa

now\\[dev\\]26
Nov 4-5, 2026San Jose Convention Center, San Jose
Register Now

Gartner IT Symposium/Xpo 2026
Oct 19-22, 2026Orlando, FL
"""


def test_markdown_fallback_for_blocked_pages():
    today = date(2026, 10, 20)
    out = extract.parse_events(SERVICENOW_MD, "http://sn", cfg(), {"label": "SN", "host": "ServiceNow"}, today,
                               markdown=True)
    by = {e.name: e for e in out}
    assert "World Forum London, United Kingdom" in by
    assert by["World Forum London, United Kingdom"].start == date(2026, 10, 20)
    assert by["now[dev]26"].city == "San Jose" and by["now[dev]26"].region == "Bay Area"
    assert all(e.companies == ["ServiceNow"] for e in out)


def test_itemlist_jsonld_and_description_tagging():
    html = ("<script type='application/ld+json'>" + json.dumps({"@type": "ItemList", "itemListElement": [
        {"@type": "ListItem", "item": {"@type": "Event", "name": "Partner Workshop", "startDate": "2027-02-10",
         "description": "Join SaaSify and AWS in Seattle", "location": {"address": {"addressLocality": "Seattle"}}}}]})
        + "</script>")
    out = extract.parse_events(html, "http://x", cfg(), {"label": "t", "require_company": True}, TODAY)
    assert [(e.name, e.companies, e.region) for e in out] == [("Partner Workshop", ["AWS"], "Seattle")]


def test_date_ranges_and_year_inference():
    t = date(2026, 10, 27)
    assert extract._range_from_match("Nov 30 – Dec 4, 2026", t) == (date(2026, 11, 30), date(2026, 12, 4))
    assert extract._range_from_match("October 25–28", t) == (date(2027, 10, 25), date(2027, 10, 28))
    assert extract._range_from_match("November 30–December 3", t) == (date(2026, 11, 30), date(2026, 12, 3))
    assert extract._range_from_match("October 20-22", t, year=2026)[0] == date(2026, 10, 20)


def test_named_date_events_flagship_and_pattern():
    page = ("<p>AWS re:Invent | Nov 30 – Dec 4, 2026 | Las Vegas, NV AWS re:Invent is a week</p>"
            "<p>October 20-22 NVIDIA GTC Berlin 2026 Connect</p><p>March 15-18 NVIDIA GTC 2027 Connect</p>")
    c, today = cfg(), date(2026, 10, 27)
    one = extract.parse_events(page, "http://aws", c, {"label": "a", "host": "AWS", "name": "AWS re:Invent", "single": True}, today)
    assert [(e.name, e.start, e.city) for e in one] == [("AWS re:Invent 2026", date(2026, 11, 30), "Las Vegas")]
    many = extract.parse_events(page, "http://nv", c, {"label": "n", "host": "NVIDIA",
                                "name_pattern": r"NVIDIA GTC(?: [A-Za-z.,' ]+?)? 20\d\d"}, today)
    assert [e.name for e in many] == ["NVIDIA GTC 2027"]  # Berlin Oct 20 2026 is already past


def test_prefix_dedupe_and_junk_names(tmp_path):
    path = tmp_path / "t.xlsx"
    a = extract.Event(name="Channel Partners Conference & Expo 2027", start=date(2027, 3, 15), city="Las Vegas", region="Other")
    b = extract.Event(name="Channel Partners Conference & Expo 2027, co-located with MSP Summit", start=date(2027, 3, 15), region="Other")
    assert workbook.merge(path, [a, b], TODAY) == (1, 1)
    html = jsonld_page([ev("Learn More & Register", "2027-03-10", "Seattle", "AWS")])
    assert extract.parse_events(html, "http://x", cfg(), {"label": "t"}, TODAY) == []


def test_small_community_events_are_dropped():
    for name, city, host in [
        ("AI Agents Seattle - August 2026", "Seattle", ""),               # no company at all
        ("Games & AI Meetup / Co-hosted by AWS", "Seattle", ""),          # meetup
        ("AWS User Group Seattle", "Seattle", "AWS User Group"),
        ("Seattle Tech Mixer 2026", "Seattle", ""),
        ("San Francisco Syntax Meetup", "San Francisco", "Syntax FM"),
        ("Claude Touch Grass Summer Social", "Seattle", "Anthropic"),
    ]:
        assert not verdict(name, city, host).keep, name


def test_default_config_is_strict_about_regions():
    shipped = load_config(ROOT / "config.yaml")
    assert shipped["keep_out_of_region"] is False
    html = jsonld_page([
        ev("Microsoft AI Tour - Toronto", "2027-02-10", "Toronto", "Microsoft"),
        ev("Microsoft AI Tour - Utrecht", "2027-02-11", "Utrecht", "Microsoft"),
        ev("Microsoft AI Tour - Seattle", "2027-02-12", "Seattle", "Microsoft"),
        ev("Microsoft AI Tour - Houston", "2027-02-13", "Houston", "Microsoft"),
        ev("Microsoft AI Tour - Mumbai", "2027-02-14", "Mumbai", "Microsoft"),
        ev("Microsoft AI Tour - New York", "2027-02-15", "New York", "Microsoft"),
        ev("Microsoft AI Tour - San Francisco", "2027-02-16", "San Francisco", "Microsoft"),
    ])
    out = extract.parse_events(html, "http://x", shipped, {"label": "t"}, TODAY)
    assert [e.region for e in out] == ["Seattle", "Texas", "India", "New York", "Bay Area"]


def test_partnership_leaders_cards_and_borrowed_city():
    html = """<div class="pl-event-box"><span class="event-date">November 18, 2026 | 8:00 am - 11:00 AM PT</span>
      <span class="event-type">In-Person Event</span><h5>Catalyst Breakfast at Microsoft Ignite</h5>
      <div class="pl-event-desc"><p>Join partnership leaders</p></div><a href="/e/1">View Event</a></div>
      <div class="pl-event-box"><span class="event-date">November 3, 2026</span><span class="event-type">In-Person Event</span>
      <h5>Catalyst Summit London</h5><a href="/e/2">View</a></div>"""
    src = {"label": "PL", "host": "Partnership Leaders", "trusted": True,
           "card_selector": {"card": ".pl-event-box", "date": ".event-date", "title": "h5", "type": ".event-type",
                             "desc": ".pl-event-desc"}}
    c = cfg(False)
    out = extract.parse_events(html, "http://pl.com/events", c, src, date(2026, 10, 28))
    assert [(e.name, e.city, e.region, e.host) for e in out] == [
        ("Catalyst Breakfast at Microsoft Ignite", "San Francisco", "Bay Area", "Partnership Leaders")]  # London dropped
    assert out[0].url == "http://pl.com/e/1"


def test_text_listing_date_first_and_city_titles_are_junk():
    text = "November 11-12, 2026\nEMEA Channel Leader Summit\nRegister\nDecember 2, 2026\nCMA Power Circles\nVirtual\nRegister\nOctober 5, 2027\nDallas, TX\n"
    src = {"label": "CMA", "host": "Channel Marketing Association", "trusted": True, "text_listing": "after"}
    out = extract.parse_events("<html><body>" + text.replace("\n", "<br>") + "</body></html>",
                               "http://cma", cfg(False), src, date(2026, 10, 28))
    assert [(e.name, e.region) for e in out] == [("CMA Power Circles", "Virtual")]


def test_year_sheets_and_earliest_date_cutoff(tmp_path):
    path = tmp_path / "t.xlsx"
    mk = lambda n, d: extract.Event(name=n, start=d, city="Seattle", region="Seattle")
    events = [mk("Too Early Summit", date(2026, 7, 31)), mk("August Summit", date(2026, 8, 1)),
              mk("Winter Summit", date(2026, 12, 9)), mk("Spring Summit 2027", date(2027, 3, 15))]
    assert workbook.merge(path, events, earliest=date(2026, 8, 1)) == (3, 3)

    from openpyxl import load_workbook
    wb = load_workbook(path)
    names = lambda sheet: [r[0] for r in wb[sheet].iter_rows(min_row=2, values_only=True)]
    assert names("2026") == ["August Summit", "Winter Summit"]
    assert names("2027") == ["Spring Summit 2027"]
    assert (tmp_path / "t_2026.csv").exists() and (tmp_path / "t_2027.csv").exists()
    # existing rows before the cutoff are purged too
    assert workbook.merge(path, [], earliest=date(2026, 12, 1)) == (0, 2)
